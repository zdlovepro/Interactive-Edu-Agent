from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.schemas.parse import PageContent, ParseResult
from app.utils.logger import logger

_FORMULA_PROG_ID_KEYWORDS = ["Equation", "MathType", "OMML"]


def _is_formula_ole(shape) -> bool:
    try:
        for element in shape._element.iter():
            tag_local = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            if tag_local == "oleObj":
                prog_id = element.get("progId", "")
                if any(keyword.lower() in prog_id.lower() for keyword in _FORMULA_PROG_ID_KEYWORDS):
                    return True
    except Exception:  # noqa: BLE001
        return False
    return False


def _is_math_shape(shape) -> bool:
    try:
        for element in shape._element.iter():
            tag_local = element.tag.split("}")[-1] if "}" in element.tag else element.tag
            if tag_local in ("oMath", "oMathPara", "m"):
                return True
    except Exception:  # noqa: BLE001
        return False
    return False


def parse_pptx(file_path: str, courseware_id: str) -> ParseResult:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件不存在：{file_path}")
    if path.suffix.lower() != ".pptx":
        raise ValueError(f"不支持的文件格式：{path.suffix}，期望 .pptx")

    logger.info("Start parsing PPTX. coursewareId=%s path=%s", courseware_id, file_path)
    presentation = Presentation(str(path))
    pages: list[PageContent] = []

    for index, slide in enumerate(presentation.slides, start=1):
        texts: list[str] = []
        image_placeholders: list[str] = []
        formula_placeholders: list[str] = []

        chart_count = 0
        diagram_count = 0
        flowchart_count = 0

        for shape in slide.shapes:
            if _is_formula_ole(shape) or _is_math_shape(shape):
                formula_placeholders.append(f"[公式:slide_{index}_formula_{len(formula_placeholders) + 1}]")
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image_placeholders.append(f"[图片:slide_{index}_img_{len(image_placeholders) + 1}]")
                continue

            # Charts / diagrams in PPTX are often not pictures.
            # python-pptx exposes chart objects via has_chart.
            try:
                if getattr(shape, "has_chart", False):
                    chart_count += 1
                    image_placeholders.append(f"[图表:slide_{index}_chart_{chart_count}]")
                    continue
            except Exception:  # noqa: BLE001
                pass

            # SmartArt / complex diagrams.
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.SMART_ART:
                    diagram_count += 1
                    image_placeholders.append(f"[示意图:slide_{index}_smartart_{diagram_count}]")
                    continue
            except Exception:  # noqa: BLE001
                pass

            # Flowcharts are commonly represented as auto-shapes (FLOWCHART_*).
            try:
                if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    auto_type = getattr(shape, "auto_shape_type", None)
                    auto_name = getattr(auto_type, "name", str(auto_type) if auto_type is not None else "")
                    if auto_name and "FLOWCHART" in str(auto_name).upper():
                        flowchart_count += 1
                        image_placeholders.append(f"[流程图:slide_{index}_flow_{flowchart_count}]")
            except Exception:  # noqa: BLE001
                pass

            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)

            if shape.has_table:
                for row in shape.table.rows:
                    row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_texts:
                        texts.append(" | ".join(row_texts))

        notes_text = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()

        pages.append(
            PageContent(
                page_index=index,
                text="\n".join(texts),
                notes=notes_text,
                image_placeholders=image_placeholders,
                formula_placeholders=formula_placeholders,
            )
        )

    result = ParseResult(
        courseware_id=courseware_id,
        file_type="pptx",
        total_pages=len(pages),
        pages=pages,
    )
    logger.info("PPTX parsing completed. coursewareId=%s totalPages=%s", courseware_id, result.total_pages)
    return result
